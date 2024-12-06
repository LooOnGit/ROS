import openpyxl
import numpy as np
################################################################################################## Biến toàn cục
file_excel = "output.xlsx" # Tên file excel để đây cho dễ sửa
################################################################################################## Hàm lưu giá trị vào file excel
def save_matrices_to_excel(w1, w2, fitness):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Train ROS2"
    sheet.cell(row=1, column=1, value="Hiden layer")
    sheet.cell(row=1, column=2, value="Output layer")
    sheet.cell(row=1, column=3, value="Fitness")
    row_index = 2
    for i in range(len(fitness)):
        matrix1_str = '\n'.join(['\t'.join(map(str, row)) for row in w1[i]])
        sheet.cell(row=row_index, column=1, value=matrix1_str)
        matrix2_str = '\n'.join(['\t'.join(map(str, row)) for row in w2[i]])
        sheet.cell(row=row_index, column=2, value=matrix2_str)
        sheet.cell(row=row_index, column=3, value=fitness[i])
        row_index += 2
    workbook.save(file_excel)
################################################################################################## Hàm đọc giá trị từ file excel
def load_from_excel():
    workbook = openpyxl.load_workbook(file_excel)
    sheet = workbook.active
    w1 = []
    w2 = []
    fitness = []
    for row in range(2, sheet.max_row + 1):
        matrix1_str = sheet.cell(row=row, column=1).value
        matrix2_str = sheet.cell(row=row, column=2).value
        fitness_value = sheet.cell(row=row, column=3).value
        if matrix1_str:
            matrix1 = [list(map(float, line.split('\t'))) for line in matrix1_str.split('\n')]
            w1.append(matrix1)
        if matrix2_str:
            matrix2 = [list(map(float, line.split('\t'))) for line in matrix2_str.split('\n')]
            w2.append(matrix2)
        if fitness_value is not None:
            fitness.append(fitness_value)
    w1 = np.array(w1)
    w2 = np.array(w2)
    return w1, w2, fitness
################################################################################################## Code mẫu lưu giá trị
# input_size = 5 # số lượng ngõ vào
# hidden_size = 2 # số neuron lớp ẩn
# output_size = 2 # số lượng ngõ ra
# ngene = 3 # số lượng cá thể
# w1 = np.random.uniform(-1,1,size=(ngene, input_size, hidden_size))
# w2 = np.random.uniform(-1,1,size=(ngene, hidden_size, output_size))
# fitness = np.random.randint(0,10000,size=(ngene))
# save_matrices_to_excel(w1, w2, fitness)
################################################################################################## Code mẫu đọc giá trị
# w1, w2, fitness = load_from_excel()
# print(w1)